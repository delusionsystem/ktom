from __future__ import annotations

import csv
import json
import logging
import os
import time
from datetime import date
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import unquote, urlencode, urlparse
from urllib.request import Request, urlopen

from .database import Database, Record, record_from_payload


LOGGER = logging.getLogger(__name__)


class DiscogsError(RuntimeError):
    """Raised when Discogs cannot provide a release result."""


class CoverCacheError(RuntimeError):
    """Raised when a release cover cannot be cached locally."""


class ReportError(RuntimeError):
    """Raised when a formatted report cannot be generated."""


class DiscogsClient:
    api_url = 'https://api.discogs.com'

    def __init__(self, token: str | None = None, timeout: float = 8.0, max_retries: int = 2) -> None:
        self.token = token or os.getenv('DISCOGS_TOKEN')
        self.timeout = timeout
        self.max_retries = max_retries

    def lookup_barcode(self, barcode: str) -> dict[str, object]:
        if not self.token:
            raise DiscogsError('Set DISCOGS_TOKEN before using Discogs lookup')
        barcode = barcode.strip()
        if not barcode:
            raise ValueError('A barcode is required')

        search_query = urlencode({'barcode': barcode, 'type': 'release', 'per_page': 1})
        search = self._request(f'{self.api_url}/database/search?{search_query}')
        results = search.get('results') or []
        if not results:
            raise DiscogsError(f'No Discogs release found for {barcode}')

        release_id = results[0].get('id')
        if not release_id:
            raise DiscogsError('Discogs returned a result without a release id')
        release = self._request(f'{self.api_url}/releases/{release_id}')
        payload = self._release_payload(release, barcode)
        try:
            prices = self._request(f'{self.api_url}/marketplace/stats/{release_id}?curr_abbr=SEK')
        except DiscogsError as error:
            LOGGER.warning('Discogs price lookup failed for release %s: %s', release_id, error)
            prices = {}
        payload.update(self._price_payload(prices))
        return payload

    def _request(self, url: str) -> dict[str, object]:
        request = Request(
            url,
            headers={
                'Accept': 'application/json',
                'User-Agent': 'KTOM/1.0',
                'Authorization': f'Discogs token={self.token}',
            },
        )
        for attempt in range(self.max_retries + 1):
            try:
                with urlopen(request, timeout=self.timeout) as response:
                    return json.loads(response.read().decode('utf-8'))
            except HTTPError as error:
                retryable = error.code == 429 or 500 <= error.code < 600
                if not retryable or attempt >= self.max_retries:
                    raise DiscogsError(f'Discogs request failed with HTTP {error.code}') from error
                retry_after = error.headers.get('Retry-After')
                try:
                    delay = min(float(retry_after), 10.0) if retry_after else 2 ** attempt
                except ValueError:
                    delay = 2 ** attempt
                LOGGER.warning('Discogs HTTP %s; retrying in %.1fs', error.code, delay)
                time.sleep(delay)
            except (URLError, TimeoutError, json.JSONDecodeError) as error:
                raise DiscogsError(f'Discogs request failed: {error}') from error

    @staticmethod
    def _release_payload(release: dict[str, object], barcode: str) -> dict[str, object]:
        artists = release.get('artists') or []
        artist = ', '.join(str(item.get('name', '')) for item in artists if item.get('name'))
        labels = release.get('labels') or []
        catalog_number = labels[0].get('catno') if labels else None
        formats = release.get('formats') or []
        media_type = formats[0].get('name') if formats else None
        return {
            'barcode': barcode,
            'discogs_release_id': release.get('id'),
            'artist': artist,
            'title': release.get('title', ''),
            'catalog_number': catalog_number,
            'media_type': media_type,
            'release_year': release.get('year'),
            'duration': DiscogsClient._duration(release.get('tracklist') or []),
            'cover_url': release.get('images', [{}])[0].get('uri') if release.get('images') else None,
            'release_type': release.get('type'),
            'credits': json.dumps(release.get('extraartists', []), ensure_ascii=False),
        }

    @staticmethod
    def _duration(tracklist: list[object]) -> str | None:
        total_seconds = 0
        found_duration = False
        for track in tracklist:
            if not isinstance(track, dict):
                continue
            duration = str(track.get('duration') or '')
            parts = duration.split(':')
            if len(parts) != 2 or not all(part.isdigit() for part in parts):
                continue
            total_seconds += int(parts[0]) * 60 + int(parts[1])
            found_duration = True
        if not found_duration:
            return None
        return f'{total_seconds // 60}:{total_seconds % 60:02d}'

    @staticmethod
    def _price_payload(stats: dict[str, object]) -> dict[str, object]:
        def value(name: str) -> float | None:
            price = stats.get(name)
            if not isinstance(price, dict):
                return None
            try:
                return float(price['value'])
            except (KeyError, TypeError, ValueError):
                return None

        return {
            'discogs_min': value('lowest_price'),
            'discogs_median': value('median_price'),
            'discogs_max': value('highest_price'),
            'price_currency': 'SEK' if stats else None,
        }


class RecordService:
    def __init__(self, database: Database | None = None, discogs: DiscogsClient | None = None) -> None:
        self.database = database or Database()
        self.discogs = discogs or DiscogsClient()

    def save(self, payload: dict[str, object]) -> Record:
        return self.database.add_record(record_from_payload(payload))

    def lookup_and_save(self, barcode: str) -> Record:
        return self.save(self.discogs.lookup_barcode(barcode))

    def search(self, query: str = '', limit: int = 100) -> list[Record]:
        return self.database.search(query, limit=limit)

    def recent(self, limit: int = 4) -> list[Record]:
        return self.database.recent(limit)

    def by_artist(self, artist: str) -> list[Record]:
        return self.database.by_artist(artist)

    def statistics(self) -> dict[str, int]:
        return self.database.statistics()

    def statistics_detail(self) -> dict[str, list[dict[str, object]]]:
        return self.database.statistics_detail()

    def get(self, record_id: int) -> Record | None:
        return self.database.get_record(record_id)

    def update(self, record_id: int, changes: dict[str, object]) -> Record:
        return self.database.update_record(record_id, changes)

    def delete(self, record_id: int) -> bool:
        return self.database.delete_record(record_id)

    def backup_database(self, destination: str | Path | None = None) -> Path:
        backup_path = destination or (
            self.database.path.parent / 'backups' / f'ktom_{date.today().isoformat()}.sqlite3'
        )
        return self.database.backup(backup_path)

    def cache_cover(self, record_id: int) -> Path:
        record = self.database.get_record(record_id)
        if not record or not record.cover_url:
            raise CoverCacheError('Record has no cover URL to cache')
        cache_dir = self.database.path.parent / 'covers'
        cache_dir.mkdir(parents=True, exist_ok=True)
        suffix = Path(unquote(urlparse(record.cover_url).path)).suffix.lower()
        if suffix not in {'.jpg', '.jpeg', '.png', '.webp'}:
            suffix = '.jpg'
        output_path = cache_dir / f'{record_id}{suffix}'
        request = Request(record.cover_url, headers={'User-Agent': 'KTOM/1.0'})
        try:
            with urlopen(request, timeout=10) as response:
                output_path.write_bytes(response.read())
        except (HTTPError, URLError, TimeoutError, OSError) as error:
            raise CoverCacheError(f'Cover download failed: {error}') from error
        self.database.update_record(record_id, {
            'cover_cached_path': str(output_path),
            'cover_cached_at': date.today().isoformat(),
        })
        return output_path

    def export_csv(self, records: list[Record] | None = None) -> Path:
        export_dir = self.database.path.parent / 'reports'
        export_dir.mkdir(parents=True, exist_ok=True)
        output_path = export_dir / f'collection_{date.today().isoformat()}.csv'
        records = records if records is not None else self.database.search()
        fields = (
            'id', 'barcode', 'discogs_release_id', 'artist', 'title', 'catalog_number', 'media_type',
            'release_year', 'duration', 'purchase_price', 'discogs_min',
            'discogs_median', 'discogs_max', 'price_currency', 'cover_url',
            'cover_cached_path', 'release_type',
            'scanned_at',
        )
        with output_path.open('w', newline='', encoding='utf-8-sig') as stream:
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            for record in records:
                writer.writerow({field: getattr(record, field) for field in fields})
        self.database.add_report(output_path, 'csv', None, len(records))
        return output_path

    def export_docx(self, records: list[Record] | None = None, query: str | None = None) -> Path:
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Cm, Mm, Pt
        except ImportError as error:
            raise ReportError('Install python-docx before generating DOCX reports') from error

        records = records if records is not None else self.database.search(query or '', limit=1000)
        if not records:
            raise ReportError('There are no records to report')

        export_dir = self.database.path.parent / 'reports'
        export_dir.mkdir(parents=True, exist_ok=True)
        output_path = export_dir / f'collection_{date.today().isoformat()}.docx'
        document = Document()
        section = document.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

        for index, record in enumerate(records):
            if index:
                document.add_page_break()
            self._add_docx_cover(document, record, WD_ALIGN_PARAGRAPH, Pt)
            document.add_page_break()
            self._add_docx_facts(document, record, Pt)

        document.save(output_path)
        self.database.add_report(output_path, 'docx', query, len(records))
        return output_path

    @staticmethod
    def _add_docx_cover(document, record: Record, alignment, point_size) -> None:
        from docx.shared import Cm

        header = document.add_paragraph('KARTOTEKSKORT - PÄRMARKIV')
        header.alignment = alignment.CENTER
        header.runs[0].font.size = point_size(9)
        title = document.add_heading(f'{record.artist} - {record.title}', level=1)
        title.alignment = alignment.CENTER
        if record.cover_cached_path and Path(record.cover_cached_path).exists():
            try:
                paragraph = document.add_paragraph()
                paragraph.alignment = alignment.CENTER
                paragraph.add_run().add_picture(record.cover_cached_path, width=Cm(11))
            except (OSError, ValueError):
                document.add_paragraph('[Cover image unavailable]').alignment = alignment.CENTER
        else:
            placeholder = document.add_paragraph('[Cover image unavailable]')
            placeholder.alignment = alignment.CENTER
        if record.cover_url:
            document.add_paragraph(f'Image URL: {record.cover_url}')

    @staticmethod
    def _add_docx_facts(document, record: Record, point_size) -> None:
        header = document.add_paragraph('FAKTABLAD - PÄRMARKIV')
        header.runs[0].font.size = point_size(9)
        document.add_heading(f'{record.artist} - {record.title}', level=1)
        fields = (
            ('ID-nummer', record.id),
            ('Katalognummer', record.catalog_number),
            ('Mediatyp', record.media_type),
            ('Utgivningsår', record.release_year),
            ('Speltid', record.duration),
            ('Inköpspris (SEK)', record.purchase_price),
            ('Discogs Min (SEK)', record.discogs_min),
            ('Discogs Median (SEK)', record.discogs_median),
            ('Discogs Max (SEK)', record.discogs_max),
            ('Skannad datum', record.scanned_at),
        )
        table = document.add_table(rows=0, cols=2)
        table.style = 'Table Grid'
        for label, value in fields:
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = str(value) if value is not None else 'Ej angivet'
